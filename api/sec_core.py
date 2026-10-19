"""
SEC data logic shared by the web API.

The functions below are copied unchanged from app.py (the Streamlit app), plus the duplicate
segment-row fix from PR #3, so the new site and the Streamlit site compute the same numbers.
app.py can't be imported directly because it runs the Streamlit page as soon as it loads.
"""
import io
import re
import pandas as pd
from docx import Document
from edgar import Company, set_identity

# SEC Compliance Identity
set_identity("Jacob Braunschweig jacob.braunschweig@gmail.com")

def is_date_or_period_column(col_name):
    """Detects whether a column represents an audited fiscal reporting period."""
    s = str(col_name).strip()
    if re.search(r"\b20[2-3][0-9]", s):
        return True
    if any(tag in s.upper() for tag in ["(FY)", "(CY)", "Q1", "Q2", "Q3", "Q4"]):
        return True
    return False

def latest_period_column(df):
    """Returns the most recent fiscal-year column (by end date), not just the first column."""
    period_cols = [c for c in df.columns if is_date_or_period_column(c)]
    if not period_cols:
        return None
    fy_cols = [c for c in period_cols if "(FY)" in str(c).upper()] or period_cols
    def end_date(c):
        m = re.search(r"\d{4}-\d{2}-\d{2}", str(c))
        return m.group(0) if m else ""
    return max(fy_cols, key=end_date)

def parse_num(val):
    """Converts an XBRL cell to float. Keeps zero, negatives and small values; returns None for blanks."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    if s.lower() in ["none", "nan", ""]:
        return None
    try:
        return float(s.replace(",", "").replace("$", "").replace("(", "-").replace(")", "").strip())
    except Exception:
        return None

def unit_kind(unit):
    """Classifies an XBRL unit name ("usd", "u_usd", "usdPerShare", "perShare", "shares", ...)."""
    u = str(unit).lower() if unit is not None and not pd.isna(unit) else "usd"
    if "share" in u:
        return "per_share" if "per" in u else "shares"
    if "usd" in u:
        return "usd"
    return "other"

def format_cell(val, unit):
    """Formats one statement value by its XBRL unit: USD in $ millions, per-share in $, shares in millions."""
    num = parse_num(val)
    if num is None:
        return ""
    kind = unit_kind(unit)
    if kind == "usd":
        return f"${num / 1_000_000:,.0f}"
    if kind == "per_share":
        return f"${num:,.2f}"
    if kind == "shares":
        return f"{num / 1_000_000:,.1f}M sh"
    return f"{num:,.2f}"

def clean_statement_df(df):
    """
    Cleans XBRL statements to match Capital IQ:
    - Retains row label and numeric fiscal date periods
    - Removes abstract category header rows that have no numeric values
    - Scales dollar values to millions ($M); per-share and share counts keep their own units
    """
    if df is None or df.empty:
        return None

    label_col = None
    for cand in ["label", "standard_concept", "concept"]:
        if cand in df.columns:
            label_col = cand
            break

    period_cols = [c for c in df.columns if is_date_or_period_column(c)]
    if not period_cols:
        return None

    work = df.copy()
    if "abstract" in work.columns:
        work = work[~work["abstract"].fillna(False).astype(bool)]

    # Drop header rows where every period value is blank or zero
    has_number = work[period_cols].apply(
        lambda row: any((parse_num(v) or 0) != 0 for v in row), axis=1
    )
    work = work[has_number]

    # Drop segment rows that just repeat their consolidated line (e.g., CAT's interest expense,
    # which is all Financial Products), so the same amount doesn't appear twice
    if "dimension" in work.columns and "concept" in work.columns:
        is_dim = work["dimension"].fillna(False).astype(bool)
        parent_values = {
            r["concept"]: tuple(r[c] for c in period_cols)
            for _, r in work[~is_dim].iterrows()
        }
        repeats = work[is_dim].apply(
            lambda r: parent_values.get(r["concept"]) == tuple(r[c] for c in period_cols), axis=1
        )
        work = work.drop(repeats[repeats].index)

    units = work["unit"] if "unit" in work.columns else pd.Series("usd", index=work.index)
    cleaned = pd.DataFrame(index=work.index)
    cleaned["Line Item"] = work[label_col] if label_col else work.index.astype(str)
    for col in period_cols:
        cleaned[col] = [format_cell(v, u) for v, u in zip(work[col], units)]
    return cleaned.set_index("Line Item")

def _norm_label(s):
    return str(s).lower().replace("’", "'").strip().rstrip(":").strip()

def extract_metric(df, concepts, labels):
    """
    Finds one consolidated line item in the latest fiscal year.
    Order: exact XBRL concepts in priority order, then exact row labels on company-specific
    (non us-gaap) concepts only, since labels on standard concepts can be misleading.
    Segment (dimensional), abstract and non-USD rows are skipped. Returns {"value", "concept", "label", "period"} or None.
    """
    if df is None or df.empty:
        return None
    period = latest_period_column(df)
    if period is None:
        return None

    rows = df
    if "dimension" in rows.columns:
        rows = rows[~rows["dimension"].fillna(False).astype(bool)]
    if "abstract" in rows.columns:
        rows = rows[~rows["abstract"].fillna(False).astype(bool)]
    if "unit" in rows.columns:
        rows = rows[rows["unit"].map(unit_kind) == "usd"]
    if rows.empty:
        return None

    # "us-gaap_Revenues" -> "Revenues"
    concept_names = rows["concept"].astype(str).str.split(r"[_:]", n=1).str[-1] if "concept" in rows.columns else None

    def first_hit(mask):
        for _, r in rows[mask].iterrows():
            n = parse_num(r[period])
            if n is not None:
                return {"value": n, "concept": str(r.get("concept", "")), "label": str(r.get("label", "")), "period": str(period)}
        return None

    if concept_names is not None:
        for c in concepts:
            hit = first_hit(concept_names == c)
            if hit:
                return hit
    if "label" in rows.columns:
        row_labels = rows["label"].map(_norm_label)
        custom = ~rows["concept"].astype(str).str.startswith("us-gaap") if "concept" in rows.columns else True
        for lbl in labels:
            hit = first_hit((row_labels == _norm_label(lbl)) & custom)
            if hit:
                return hit
    return None

def compute_ratios(raw_inc, raw_bal):
    """Pulls the ratio inputs from the latest 10-K and computes the four ratios. Returns (ratios, inputs)."""
    revenue = extract_metric(raw_inc, ["Revenues", "RevenuesNetOfInterestExpense"], [])
    if revenue is None:
        # Banks without a total revenue line: net interest income + noninterest income
        nii = extract_metric(raw_inc, ["InterestIncomeExpenseNet"], [])
        nonint = extract_metric(raw_inc, ["NoninterestIncome"], [])
        if nii and nonint:
            revenue = {"value": nii["value"] + nonint["value"], "concept": "Derived: Net interest income + Noninterest income",
                       "label": "Net interest income + Total noninterest income", "period": nii["period"]}
    if revenue is None:
        revenue = extract_metric(
            raw_inc,
            ["RevenueFromContractWithCustomerExcludingAssessedTax", "RevenueFromContractWithCustomerIncludingAssessedTax", "SalesRevenueNet"],
            ["Total sales and revenues", "Total net sales", "Total revenues", "Total net revenue", "Net sales", "Revenues", "Revenue"]
        )

    inputs = {
        "Revenue": revenue,
        "Cost of Revenue": extract_metric(
            raw_inc,
            ["CostOfGoodsAndServicesSold", "CostOfRevenue", "CostOfGoodsSold", "CostOfGoodsAndServiceExcludingDepreciationDepletionAndAmortization"],
            ["Cost of goods sold", "Cost of sales", "Cost of products sold", "Cost of revenue", "Total cost of revenue", "Total cost of sales"]
        ),
        "Gross Profit": extract_metric(raw_inc, ["GrossProfit"], ["Gross profit", "Gross margin"]),
        "Operating Income": extract_metric(
            raw_inc,
            ["OperatingIncomeLoss"],
            ["Operating profit", "Operating income", "Operating income (loss)", "Income from operations"]
        ),
        "Current Assets": extract_metric(raw_bal, ["AssetsCurrent"], ["Total current assets"]),
        "Current Liabilities": extract_metric(raw_bal, ["LiabilitiesCurrent"], ["Total current liabilities"]),
        "Total Liabilities": extract_metric(raw_bal, ["Liabilities"], ["Total liabilities"]),
        "Shareholders' Equity": extract_metric(
            raw_bal,
            ["StockholdersEquity", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"],
            ["Total shareholders' equity", "Total stockholders' equity", "Total equity"]
        ),
    }

    # Derived inputs are labeled as such so every number still traces to filed line items
    rev, cogs = inputs["Revenue"], inputs["Cost of Revenue"]
    if inputs["Gross Profit"] is None and rev and cogs:
        inputs["Gross Profit"] = {"value": rev["value"] - cogs["value"], "concept": "Derived: Revenue - Cost of Revenue",
                                  "label": "Revenue - Cost of Revenue", "period": rev["period"]}
    if inputs["Total Liabilities"] is None:
        liab_and_eq = extract_metric(raw_bal, ["LiabilitiesAndStockholdersEquity"], ["Total liabilities and shareholders' equity", "Total liabilities and stockholders' equity", "Total liabilities and equity"])
        total_eq = extract_metric(
            raw_bal,
            ["StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest", "StockholdersEquity"],
            ["Total equity", "Total shareholders' equity", "Total stockholders' equity"]
        )
        if liab_and_eq and total_eq:
            inputs["Total Liabilities"] = {"value": liab_and_eq["value"] - total_eq["value"],
                                           "concept": "Derived: Total liabilities & equity - Total equity",
                                           "label": "Total liabilities & equity - Total equity", "period": liab_and_eq["period"]}

    def v(name):
        return inputs[name]["value"] if inputs[name] else None

    def pct(num, den):
        return f"{num / den * 100:.1f}%" if num is not None and den is not None and den > 0 else "N/A"

    def times(num, den):
        return f"{num / den:.2f}x" if num is not None and den is not None and den > 0 else "N/A"

    equity = v("Shareholders' Equity")
    ratios = {
        "Gross Margin": pct(v("Gross Profit"), v("Revenue")),
        "Operating Margin": pct(v("Operating Income"), v("Revenue")),
        "Current Ratio": times(v("Current Assets"), v("Current Liabilities")),
        # Negative equity makes the ratio meaningless, so say so instead of showing a negative multiple
        "Total Liabilities-to-Equity": "N/M (negative equity)" if equity is not None and equity < 0
                                       else times(v("Total Liabilities"), equity),
    }
    return ratios, inputs

def clean_governance_text(raw_text):
    """
    Cleans raw SEC proxy text (works for any filer):
    - Removes underscores, horizontal lines, dashes, and boilerplate rules
    - Removes running "Proxy Statement" page headers and mailing addresses
    - Collapses whitespace
    """
    if not raw_text:
        return ""
    text = raw_text
    # Drop invisible characters and unify dashes/apostrophes so keyword searches match
    text = re.sub(r"[\u00ad\u200b\u200c\u200d\ufeff]", "", text)
    text = re.sub(r"[\u2010\u2011\u2012\u2013]", "-", text)
    text = text.replace("\u2019", "'")
    # Remove lines of underscores, dashes, equals signs
    text = re.sub(r"_{2,}", " ", text)
    text = re.sub(r"-{3,}", " ", text)
    text = re.sub(r"={3,}", " ", text)
    # Remove running SEC headers/footers and addresses
    text = re.sub(r"\b\d{4}\s+Proxy Statement\s+\d{1,3}\b", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"Proxy Statement\s*\d{4}\s*\d{1,3}", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"c/o Corporate Secretary.{0,200}?\d{5}", " ", text, flags=re.IGNORECASE)
    # Collapse whitespace
    text = re.sub(r"\s+", " ", text).strip()
    return text

# Keyword checks for the governance tiles. A hit means the proxy discusses the topic; the matching
# sentence is shown as evidence so the reader can confirm what the filing actually says.
GOVERNANCE_CHECKS = {
    "Lead Independent Director": [r"lead independent director", r"\blead director\b", r"presiding (?:independent )?director"],
    "Clawback Policy": [r"clawback", r"(?:incentive )?compensation recovery policy"],
    "Anti-Hedging Policy": [r"anti-hedging", r"hedging (?:and|or) pledging", r"prohibit\w*[^.]{0,200}hedg"],
    "Say-on-Pay Vote": [r"advisory vote to approve (?:the )?(?:compensation of (?:our )?named executive officers|(?:named )?executive (?:officer )?compensation)",
                        r"say[- ]on[- ]pay"],
}

def html_to_plain_text(html):
    """Strips tags from an HTML document, keeping block breaks as newlines."""
    from bs4 import BeautifulSoup
    if isinstance(html, bytes):
        html = html.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    # Inline XBRL proxies hide header facts in ix:header; drop them so they don't pollute the text
    for tag in soup.find_all(["ix:header"]):
        tag.decompose()
    return soup.get_text("\n")

def get_proxy_text(filing):
    """
    Returns (text, source) for a DEF 14A. Tries edgartools' text() first; if that fails or comes back
    too short to be a full proxy, reads every HTML document in the filing directly, since some
    proxies are split across several files or don't parse with the default extractor.
    """
    text, source = "", "none"
    try:
        text = filing.text() or ""
        source = "primary document"
    except Exception:
        text = ""
    if len(text) >= 50000:
        return text, source

    parts = []
    try:
        for att in filing.attachments:
            if not att.is_html() or att.is_report() or att.empty:
                continue
            content = att.download()
            if content:
                parts.append(html_to_plain_text(content))
    except Exception:
        pass
    combined = "\n\n".join(parts)
    if len(combined) > len(text):
        return combined, f"{len(parts)} HTML document(s) in the filing"
    return text, source

def find_evidence(text, patterns):
    """Returns the sentence around the first pattern match in the proxy, or None."""
    for p in patterns:
        m = re.search(p, text, flags=re.IGNORECASE)
        if m:
            start = max(text.rfind(". ", 0, m.start()) + 2, m.start() - 300, 0)
            end = text.find(". ", m.end())
            end = min(end + 1 if end != -1 else len(text), m.end() + 300)
            return text[start:end].strip()
    return None

def build_word_memo(ticker_sym, company_nm, ratios_dict, gov_dict, memo_body):
    """Generates a structured Word (.docx) Institutional Research Memo."""
    doc = Document()
    doc.add_heading(f"Institutional Research & Corporate Governance Brief", level=0)
    doc.add_paragraph(f"Company: {company_nm} ({ticker_sym}) | Source: Audited SEC EDGAR 10-K & DEF 14A")

    doc.add_heading("1. Executive Summary & Audited Ratios", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Financial Ratio"
    hdr_cells[1].text = "Audited Metric Value"
    for metric_name, val in ratios_dict.items():
        row_cells = table.add_row().cells
        row_cells[0].text = metric_name
        row_cells[1].text = str(val)

    doc.add_heading("2. Qualitative Corporate Governance Disclosures (DEF 14A)", level=1)
    for section_title, section_content in gov_dict.items():
        doc.add_heading(section_title, level=2)
        doc.add_paragraph(section_content[:1500] + ("..." if len(section_content) > 1500 else ""))

    doc.add_heading("3. Institutional Equity Research Synthesis", level=1)
    doc.add_paragraph(memo_body)

    file_buffer = io.BytesIO()
    doc.save(file_buffer)
    return file_buffer.getvalue()

def statement_frames(statement):
    """Returns (display_df, raw_df) for one statement, or (None, None) if the filing lacks it."""
    if statement is None:
        return None, None
    # standard=False keeps the filer's own line-item labels instead of edgartools' standardized ones
    display_df = statement.to_dataframe(view="standard", include_unit=True, standard=False)
    raw_df = statement.to_dataframe(presentation=False, include_unit=True, standard=False)
    return display_df, raw_df

SECTION_TARGETS = {
    "Board Leadership & Lead Independent Director Structure": [
        "board leadership structure", "lead independent director", "director independence"
    ],
    "Executive Compensation & Pay-for-Performance (CD&A)": [
        "compensation discussion and analysis", "executive compensation program",
        "compensation philosophy", "summary compensation table"
    ],
    "Annual Meeting Agenda & Shareholder Voting Items": [
        "matters to be voted on", "proposals to be voted on", "election of directors",
        "shareholder proposals", "stockholder proposals"
    ],
    "Clawback Policies, Anti-Hedging & Risk Oversight": [
        "clawback policy", "compensation recovery policy",
        "anti-hedging", "hedging and pledging", "risk oversight"
    ]
}

def statement_json(df):
    """Turns a cleaned statement table into {"periods": [...], "rows": [{"label", "values"}]} for the web page."""
    if df is None or df.empty:
        return None
    return {
        "periods": [str(c) for c in df.columns],
        "rows": [{"label": str(label), "values": [str(v) for v in row]} for label, row in df.iterrows()],
    }

def section_passage(text_body):
    """Shows the passage as filed, cut at the last full sentence (same rule as app.py after PR #3)."""
    passage = text_body[:2000]
    last_stop = passage.rfind(". ")
    if len(text_body) > 2000 and last_stop > 500:
        passage = passage[:last_stop + 1]
    return passage

def get_company_report(ticker_symbol):
    """Everything the web page shows for one ticker, as plain JSON-ready data. Same steps as get_sec_data in app.py."""
    company = Company(ticker_symbol)

    financials = company.get_financials()
    disp_income, raw_income = statement_frames(financials.income_statement() if financials else None)
    disp_balance, raw_balance = statement_frames(financials.balance_sheet() if financials else None)
    disp_cashflow, _ = statement_frames(financials.cash_flow_statement() if financials else None)

    ratios, inputs = compute_ratios(raw_income, raw_balance)

    proxy_filings = company.get_filings(form="DEF 14A")
    tenk_filings = company.get_filings(form="10-K")

    governance = {"checks": {}, "sections": {}, "proxy_url": "", "proxy_date": "", "proxy_chars": 0, "proxy_source": ""}
    if proxy_filings:
        latest_proxy = proxy_filings[0]
        proxy_raw, proxy_source = get_proxy_text(latest_proxy)
        proxy_flat = clean_governance_text(proxy_raw)
        governance.update(proxy_url=latest_proxy.url, proxy_date=str(latest_proxy.filing_date),
                          proxy_chars=len(proxy_raw), proxy_source=proxy_source)
        for check, patterns in GOVERNANCE_CHECKS.items():
            governance["checks"][check] = find_evidence(proxy_flat, patterns)

        # Prefer matches after the first 15,000 characters to bypass the table of contents and cover page
        proxy_lower = proxy_flat.lower()
        body_start = 15000 if len(proxy_flat) > 25000 else 0
        for section_title, keywords in SECTION_TARGETS.items():
            found_pos = -1
            for kw in keywords:
                pos = proxy_lower.find(kw, body_start)
                if pos == -1:
                    pos = proxy_lower.find(kw)
                if pos != -1:
                    found_pos = pos
                    break
            governance["sections"][section_title] = (
                section_passage(proxy_flat[found_pos : found_pos + 3000]) if found_pos != -1
                else "No matching section heading was found in the proxy text."
            )

    return {
        "ticker": ticker_symbol,
        "name": company.name,
        "tenk_url": tenk_filings[0].url if tenk_filings else "",
        "statements": {
            "income": statement_json(clean_statement_df(disp_income)),
            "balance": statement_json(clean_statement_df(disp_balance)),
            "cashflow": statement_json(clean_statement_df(disp_cashflow)),
        },
        "ratios": ratios,
        # Values in $ millions, with the filed line item and XBRL concept each one came from
        "inputs": {
            name: ({"value_m": hit["value"] / 1_000_000, "label": hit["label"], "concept": hit["concept"], "period": hit["period"]}
                   if hit else None)
            for name, hit in inputs.items()
        },
        "governance": governance,
    }
