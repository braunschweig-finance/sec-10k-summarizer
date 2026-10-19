"""
Excel and PDF downloads of one company's numbers. Both are built from the same JSON the page shows,
so every figure in the files matches the site and keeps its XBRL concept and filing next to it.
"""
import io
from datetime import date

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
M = '#,##0;(#,##0)'          # $ millions
M1 = '#,##0.0;(#,##0.0)'     # $ millions, one decimal (CEO pay)
PCT = '0.00%'


def _sheet(wb, title, header, rows, formats=None, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append(r)
    for col, fmt in (formats or {}).items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _pct(v):
    """Percent value (e.g. 1.18) to an Excel fraction (0.0118), keeping blanks blank."""
    return v / 100 if v is not None else None


def build_xlsx(data):
    r, pe, tr = data["report"], data.get("pay_efficiency"), data.get("trends")
    wb = Workbook()
    wb.remove(wb.active)

    summary = [
        ["Company", r["name"]],
        ["Ticker", r["ticker"]],
        ["10-K", r.get("tenk_url", "")],
        ["Exported", date.today().isoformat()],
        ["Source", "SEC EDGAR filings (10-K XBRL, DEF 14A Pay Versus Performance XBRL). No AI-generated numbers."],
    ]
    summary += [[f"{name}", value] for name, value in r["ratios"].items()]
    if pe and pe.get("available"):
        summary.append([f"Pay Efficiency (FY{pe['fy_end'][:4]})", pe["value"]])
    summary += [["Note", n] for n in data.get("notes", [])]
    _sheet(wb, "Summary", ["Item", "Value"], summary, widths=[34, 90])

    _sheet(
        wb, "Ratio inputs",
        ["Input", "Value ($M)", "Filed line item", "XBRL concept", "Period"],
        [[name, h["value_m"] if h else None, h["label"] if h else "Not reported", h["concept"] if h else "", h["period"] if h else ""]
         for name, h in r["inputs"].items()],
        {"B": M}, [24, 14, 44, 56, 24],
    )

    if pe and pe.get("available"):
        rows = [[f"FY{y['fy_end'][:4]}", y["fy_end"], y["ceo_pay_m"], y["op_income_m"], _pct(y["pct"]), y["value"],
                 y["op_concept"], y["tenk_accn"]] for y in pe.get("history", [])]
        ws = _sheet(
            wb, "Pay Efficiency",
            ["Fiscal year", "Year end", "CEO pay actually paid ($M)", "Operating income ($M)", "Pay Efficiency", "As shown",
             "Denominator concept", "10-K accession"],
            rows, {"C": M1, "D": M, "E": PCT}, [12, 12, 26, 22, 15, 22, 52, 24],
        )
        ws.append([])
        ws.append(["Pay concept", pe["pay_concept"]])
        ws.append(["Proxy (DEF 14A)", pe["proxy_url"]])
        ws.append(["Proxy accession", pe["proxy_accn"]])
        for f in pe.get("flags", []):
            ws.append(["Flag", f])
        peers = pe.get("peers")
        if peers:
            ws.append([])
            ws.append(["Peer ranking", f"Lower than {peers['lower_than_pct']:.0f}% of {peers['count']} peers (fiscal {peers['peer_year']})"])
            ws.append(["Peer median", _pct(peers["median"])])
            ws.cell(ws.max_row, 2).number_format = PCT
            ws.append(["Peer 25th-75th percentile", f"{peers['p25']:.2f}% to {peers['p75']:.2f}%"])
            ws.append(["Peer sample", peers["source"]])
            for n in peers.get("notes", []):
                ws.append(["Note", n])

    if tr and tr.get("years"):
        names = ["Revenue", "Cost of Revenue", "Gross Profit", "Operating Income", "Net Income"]
        rows = []
        for y in tr["years"]:
            lines = y["lines"]
            rows.append(
                [y["fiscal_year"], y["fy_end"]]
                + [lines[n]["value_m"] if lines[n] else None for n in names]
                + [_pct(y["margins"][m]) for m in ["Gross Margin", "Operating Margin", "Net Margin"]]
                + [lines["Revenue"]["accn"] if lines["Revenue"] else ""]
            )
        ws = _sheet(
            wb, "Trends",
            ["Fiscal year", "Year end"] + [f"{n} ($M)" for n in names] + ["Gross margin", "Operating margin", "Net margin", "10-K accession"],
            rows, {"C": M, "D": M, "E": M, "F": M, "G": M, "H": PCT, "I": PCT, "J": PCT},
            [12, 12, 15, 19, 16, 21, 15, 14, 17, 12, 24],
        )
        ws.append([])
        ws.append(["Each year's figures are as first reported in that year's 10-K (SEC companyfacts). Concepts used:"])
        for y in tr["years"]:
            used = ", ".join(f"{n}: {l['concept']}" for n, l in y["lines"].items() if l)
            ws.append([y["fiscal_year"], used])

    for key, title in [("income", "Income statement"), ("balance", "Balance sheet"), ("cashflow", "Cash flow")]:
        st = r["statements"].get(key)
        if st:
            _sheet(wb, title, ["Line item"] + st["periods"], [[row["label"]] + row["values"] for row in st["rows"]],
                   widths=[60] + [18] * len(st["periods"]))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _latin1(text):
    """The PDF's built-in fonts only cover Latin-1, so swap the few characters filings use outside it."""
    text = str(text)
    for a, b in {"’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-", "…": "...", " ": " "}.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


class _Pdf(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(120)
        self.cell(0, 6, _latin1(f"Source: SEC EDGAR filings. Generated {date.today().isoformat()} by the SEC Financial & Corporate Governance Terminal."), align="L")
        self.cell(0, 6, f"Page {self.page_no()}", align="R")


def _table(pdf, header, rows, widths, align=None):
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(31, 41, 55)
    pdf.set_text_color(255)
    for j, (h, w) in enumerate(zip(header, widths)):
        pdf.cell(w, 6, _latin1(h), border=0, fill=True, align=(align[j] if align else "L"))
    pdf.ln()
    pdf.set_font("Helvetica", size=8)
    pdf.set_text_color(20)
    for i, row in enumerate(rows):
        pdf.set_fill_color(243, 244, 246) if i % 2 else pdf.set_fill_color(255, 255, 255)
        for j, (v, w) in enumerate(zip(row, widths)):
            text = _latin1(v)
            while pdf.get_string_width(text) > w - 1 and len(text) > 3:
                text = text[:-4] + "..."
            pdf.cell(w, 5.5, text, fill=True, align=(align[j] if align else "L"))
        pdf.ln()
    pdf.ln(3)


def _h(pdf, text):
    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(20)
    pdf.cell(0, 8, _latin1(text), new_x="LMARGIN", new_y="NEXT")


def _p(pdf, text, size=8.5):
    pdf.set_font("Helvetica", size=size)
    pdf.set_text_color(70)
    pdf.multi_cell(0, 4.5, _latin1(text), align="L", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)


def _m(v):
    return f"{v:,.0f}" if v is not None else "-"


def _pc(v):
    return f"{v:.1f}%" if v is not None else "N/A"


def build_pdf(data):
    r, pe, tr = data["report"], data.get("pay_efficiency"), data.get("trends")
    pdf = _Pdf(format="Letter")
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, _latin1(f"{r['name']} ({r['ticker']})"), new_x="LMARGIN", new_y="NEXT")
    _p(pdf, f"Financial ratios, Pay Efficiency and multi-year trends from SEC EDGAR filings. 10-K: {r.get('tenk_url', '')}")
    for n in data.get("notes", []):
        _p(pdf, n)

    _h(pdf, "Ratios (latest 10-K)")
    _table(pdf, ["Ratio", "Value"], [[k, v] for k, v in r["ratios"].items()], [70, 40], ["L", "R"])
    _table(
        pdf, ["Input", "$M", "Filed line item", "XBRL concept", "Period"],
        [[k, _m(h["value_m"]) if h else "Not reported", h["label"] if h else "", h["concept"] if h else "", h["period"] if h else ""]
         for k, h in r["inputs"].items()],
        [32, 20, 50, 66, 28], ["L", "R", "L", "L", "L"],
    )

    if pe and pe.get("available"):
        _h(pdf, f"Pay Efficiency (FY{pe['fy_end'][:4]}): {pe['value']}")
        _p(pdf, "CEO compensation actually paid (DEF 14A Pay Versus Performance table, ecd:PeoActuallyPaidCompAmt) divided by "
                f"operating income from the 10-K. Lower means the company keeps more operating profit per dollar of CEO pay. Proxy: {pe['proxy_url']}")
        _table(
            pdf, ["Fiscal year", "CEO pay actually paid ($M)", "Operating income ($M)", "Pay Efficiency", "10-K accession"],
            [[f"FY{y['fy_end'][:4]}", f"{y['ceo_pay_m']:,.1f}", _m(y["op_income_m"]), y["value"], y["tenk_accn"]] for y in pe.get("history", [])],
            [24, 42, 38, 30, 46], ["L", "R", "R", "R", "L"],
        )
        peers = pe.get("peers")
        if peers:
            _p(pdf, f"Peer ranking: lower than {peers['lower_than_pct']:.0f}% of {peers['count']} peers in fiscal {peers['peer_year']} "
                    f"(peer median {peers['median']:.2f}%, middle half {peers['p25']:.2f}% to {peers['p75']:.2f}%). Peers: {peers['source']}. "
                    + " ".join(peers.get("notes", [])))
        for f in pe.get("flags", []):
            _p(pdf, f"Flag: {f}")

    if tr and tr.get("years"):
        _h(pdf, "Multi-year trends ($ millions)")
        _table(
            pdf, ["Fiscal year", "Revenue", "Gross profit", "Operating inc.", "Net income", "Gross m.", "Op. m.", "Net m."],
            [[str(y["fiscal_year"])] + [_m(y["lines"][n]["value_m"]) if y["lines"][n] else "-"
                                        for n in ["Revenue", "Gross Profit", "Operating Income", "Net Income"]]
             + [_pc(y["margins"][m]) for m in ["Gross Margin", "Operating Margin", "Net Margin"]] for y in tr["years"]],
            [22, 26, 26, 26, 26, 22, 22, 22], ["L"] + ["R"] * 7,
        )
        _p(pdf, "Each year is as first reported in that year's 10-K (SEC companyfacts). Accessions: "
                + "; ".join(f"FY{y['fiscal_year']} {y['lines']['Revenue']['accn']}" for y in tr["years"] if y["lines"]["Revenue"]))

    return bytes(pdf.output())
