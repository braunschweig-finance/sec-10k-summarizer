"""
Offline tests for trends, Pay Efficiency history, peer ranking and the export files.
SEC is replaced by small hand-made companyfacts and Pay Versus Performance files, so the expected
numbers can be checked by hand. Run from api/:  python -m pytest tests
"""
import io
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from openpyxl import load_workbook

import export
import peers
import sec_core


def fact(val, start, end, accn, filed, form="10-K"):
    return {"val": val, "start": start, "end": end, "accn": accn, "filed": filed, "form": form}


def year_facts(year, val, restated=None):
    """One fiscal year (calendar) as first filed in its own 10-K, plus an optional later restatement."""
    s, e = f"{year}-01-01", f"{year}-12-31"
    rows = [fact(val, s, e, f"0000018230-{str(year + 1)[2:]}-000008", f"{year + 1}-02-15")]
    if restated is not None:
        rows.append(fact(restated, s, e, f"0000018230-{str(year + 2)[2:]}-000008", f"{year + 2}-02-15"))
    # A quarter and a 10-Q, both of which must be ignored
    rows.append(fact(val / 4, f"{year}-10-01", e, "q", f"{year + 1}-02-15"))
    rows.append(fact(val * 9, s, e, "10q", f"{year}-11-01", form="10-Q"))
    return rows


GAAP = {
    "Revenues": {"units": {"USD": sum((year_facts(y, v) for y, v in
                                       [(2019, 50e9), (2020, 40e9), (2021, 50e9), (2022, 60e9), (2023, 67e9), (2024, 64e9), (2025, 68e9)]), [])}},
    "CostOfGoodsAndServicesSold": {"units": {"USD": sum((year_facts(y, v) for y, v in
                                                         [(2021, 35e9), (2022, 42e9), (2023, 44e9), (2024, 42e9), (2025, 45e9)]), [])}},
    # Operating income restated later for 2024: the trend must keep the first-reported 13e9
    "OperatingIncomeLoss": {"units": {"USD": sum((year_facts(y, v, 13.5e9 if y == 2024 else None) for y, v in
                                                  [(2021, 6e9), (2022, 8e9), (2023, 13e9), (2024, 13e9), (2025, 11e9)]), [])}},
    "NetIncomeLoss": {"units": {"USD": sum((year_facts(y, v) for y, v in
                                            [(2021, 6.5e9), (2022, 6.7e9), (2023, 10e9), (2024, 10.8e9), (2025, 8.9e9)]), [])}},
}


def pvp_xml(years):
    """A minimal Pay Versus Performance instance: TSR per year, and CEO pay actually paid per CEO."""
    ctx, facts = [], []
    for i, (end, ceos) in enumerate(years):
        ctx.append(f'<xbrli:context id="c{i}"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier></xbrli:entity>'
                   f'<xbrli:period><xbrli:startDate>{end[:4]}-01-01</xbrli:startDate><xbrli:endDate>{end}</xbrli:endDate></xbrli:period></xbrli:context>')
        facts.append(f'<ecd:TotalShareholderRtnAmt contextRef="c{i}">100</ecd:TotalShareholderRtnAmt>')
        for j, (who, cap) in enumerate(ceos):
            ctx.append(f'<xbrli:context id="c{i}_{j}"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier>'
                       f'<xbrli:segment><xbrldi:explicitMember dimension="ecd:IndividualAxis">{who}</xbrldi:explicitMember></xbrli:segment></xbrli:entity>'
                       f'<xbrli:period><xbrli:startDate>{end[:4]}-01-01</xbrli:startDate><xbrli:endDate>{end}</xbrli:endDate></xbrli:period></xbrli:context>')
            facts.append(f'<ecd:PeoActuallyPaidCompAmt contextRef="c{i}_{j}">{cap}</ecd:PeoActuallyPaidCompAmt>')
            facts.append(f'<ecd:PeoTotalCompAmt contextRef="c{i}_{j}">{cap / 2}</ecd:PeoTotalCompAmt>')
    return ('<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:xbrldi="http://xbrl.org/2006/xbrldi" '
            'xmlns:ecd="http://xbrl.sec.gov/ecd/2024">' + "".join(ctx) + "".join(facts) + "</xbrli:xbrl>").encode()


PVP = pvp_xml([
    ("2023-12-31", [("CAT:Mr.UmplebyMember", 26e6)]),
    ("2024-12-31", [("CAT:Mr.UmplebyMember", 39e6)]),
    # Two CEOs in the latest year (CAT's real FY2025 amounts): added together
    ("2025-12-31", [("CAT:Mr.UmplebyMember", 82_309_936), ("CAT:Mr.CreedMember", 49_423_389)]),
])


@pytest.fixture
def fake_sec(monkeypatch):
    proxy = SimpleNamespace(accession_no="0001308179-26-000358", url="https://www.sec.gov/proxy", filing_date="2026-04-30")
    company = SimpleNamespace(cik=18230, name="Caterpillar Inc.", get_filings=lambda form, trigger_full_load=True: [proxy])
    monkeypatch.setattr(sec_core, "Company", lambda t: company)
    monkeypatch.setattr(sec_core, "companyfacts", lambda cik: GAAP)
    monkeypatch.setattr(sec_core, "instance_url", lambda cik, accn: "https://x/proxy_htm.xml")
    monkeypatch.setattr(sec_core, "sec_fetch", lambda url: PVP)
    monkeypatch.setattr(sec_core.prices, "year_end_closes", lambda t, ends: {"available": False, "reason": "offline test", "source": ""})


def test_trends_take_last_five_years_as_first_reported(fake_sec):
    t = sec_core.get_trends("CAT")
    assert [y["fiscal_year"] for y in t["years"]] == [2021, 2022, 2023, 2024, 2025]
    y24 = t["years"][3]
    assert y24["lines"]["Operating Income"]["value_m"] == 13_000            # not the 13.5B restatement
    assert y24["lines"]["Operating Income"]["accn"] == "0000018230-25-000008"
    assert y24["lines"]["Revenue"]["concept"] == "us-gaap:Revenues"
    # No GrossProfit tag: derived from revenue - cost of revenue, and labeled as derived
    assert y24["lines"]["Gross Profit"]["value_m"] == 64_000 - 42_000
    assert y24["lines"]["Gross Profit"]["concept"].startswith("Derived")
    assert y24["margins"]["Gross Margin"] == pytest.approx(22 / 64 * 100)    # 34.375%
    assert y24["margins"]["Operating Margin"] == pytest.approx(13 / 64 * 100)  # 20.3125%
    assert y24["margins"]["Net Margin"] == pytest.approx(10.8 / 64 * 100)    # 16.875%


def test_pay_efficiency_history_and_latest(fake_sec):
    pe = sec_core.get_pay_efficiency("CAT")
    assert pe["fy_end"] == "2025-12-31"
    # (82,309,936 + 49,423,389) / 11,000,000,000 = 1.1976%
    assert pe["ceo_pay_m"] == pytest.approx(131.733325)
    assert pe["value"] == "1.20%"
    assert pe["ceo_count"] == 2 and pe["ceos"] == ["Mr. Umpleby", "Mr. Creed"]
    assert [h["fy_end"] for h in pe["history"]] == ["2023-12-31", "2024-12-31", "2025-12-31"]
    assert pe["history"][0]["pct"] == pytest.approx(26 / 13_000 * 100)   # 0.20%
    assert pe["history"][1]["pct"] == pytest.approx(39 / 13_000 * 100)   # 0.30%, first-reported 13.0B
    assert pe["peers"]["peer_year"] == 2025 and pe["peers"]["count"] == 120


def test_peer_rank_counts_by_hand():
    group = [p["pct"] for p in peers.PEERS[2025] if p["ticker"] != "CAT"]
    r = peers.rank("CAT", "2025-12-31", 1.18, "operating")
    assert r["higher_count"] == sum(1 for v in group if v > 1.18)
    assert r["lower_than_pct"] == pytest.approx(r["higher_count"] / len(group) * 100)
    s = sorted(group)
    assert r["median"] == pytest.approx((s[59] + s[60]) / 2)               # 120 values: average of the middle two


def test_peer_rank_leaves_company_out_and_falls_back_a_year():
    a = peers.rank("A", "2025-10-31", 1.0, "operating")
    assert a["count"] == len(peers.PEERS[2025]) - 1                    # Agilent is in the sample
    later = peers.rank("CAT", "2026-12-31", 1.0, "pretax")
    assert later["peer_year"] == 2025 and len(later["notes"]) == 2
    assert peers.rank("CAT", "2025-12-31", None, "") is None
    assert peers.fiscal_year("2026-01-03") == 2025                     # 52/53-week year ending in early January


def test_peer_file_matches_its_source_math():
    # Agilent FY2022 from the research data: 10,456,854 / 1,618,000,000
    a22 = next(p for p in peers.PEERS[2022] if p["ticker"] == "A")
    assert a22["pct"] == pytest.approx(10_456_854 / 1_618_000_000 * 100)
    assert a22["pvp_accn"] == "0000950170-24-010096"


def export_input():
    pe = sec_core.get_pay_efficiency("CAT")
    tr = sec_core.get_trends("CAT")
    report = {
        "ticker": "CAT", "name": "Caterpillar Inc.’s test", "tenk_url": "https://www.sec.gov/10k",
        "statements": {"income": {"periods": ["2025-12-31 (FY)"], "rows": [{"label": "Revenue", "values": ["$68,000"]}]}, "balance": None, "cashflow": None},
        "ratios": {"Gross Margin": "33.8%", "Operating Margin": "16.2%", "Current Ratio": "1.44x", "Total Liabilities-to-Equity": "3.62x"},
        "inputs": {"Revenue": {"value_m": 68000.0, "label": "Total sales and revenues", "concept": "us-gaap_Revenues", "period": "2025-12-31 (FY)"},
                   "Gross Profit": None},
    }
    return {"report": report, "pay_efficiency": pe, "trends": tr, "notes": []}


def test_xlsx_export_has_every_sheet_and_numbers(fake_sec):
    wb = load_workbook(io.BytesIO(export.build_xlsx(export_input())))
    assert wb.sheetnames == ["Summary", "Ratio inputs", "Pay Efficiency", "Trends", "Income statement"]
    pe = wb["Pay Efficiency"]
    assert pe["E4"].value == pytest.approx(131.733325 / 11_000)   # stored as a fraction, shown as 1.20%
    tr = wb["Trends"]
    assert tr["A6"].value == 2025 and tr["C6"].value == 68_000 and tr["I6"].value == pytest.approx(11 / 68)
    assert wb["Ratio inputs"]["C3"].value == "Not reported"


def test_pdf_export_builds(fake_sec):
    data = export.build_pdf(export_input())
    assert data[:5] == b"%PDF-" and len(data) > 2000


def test_export_skips_missing_parts():
    d = {"report": {"ticker": "JPM", "name": "JPMorgan", "tenk_url": "", "statements": {}, "ratios": {"Gross Margin": "N/A"}, "inputs": {}},
         "pay_efficiency": None, "trends": None, "notes": ["Trends left out: timed out"]}
    assert load_workbook(io.BytesIO(export.build_xlsx(d))).sheetnames == ["Summary", "Ratio inputs"]
    assert export.build_pdf(d)[:5] == b"%PDF-"


def test_api_routes(fake_sec, monkeypatch):
    from fastapi.testclient import TestClient
    import main
    report = export_input()["report"]
    monkeypatch.setattr(sec_core, "get_financial_report", lambda t: report)
    for f in [main.cached_financials, main.cached_pay_efficiency, main.cached_trends]:
        f.cache_clear()
    client = TestClient(main.app)
    assert [y["fiscal_year"] for y in client.get("/api/trends/cat").json()["years"]][-1] == 2025
    x = client.get("/api/export/CAT.xlsx")
    assert x.status_code == 200 and x.headers["content-disposition"] == 'attachment; filename="CAT_SEC_Terminal.xlsx"'
    p = client.get("/api/export/CAT.pdf")
    assert p.status_code == 200 and p.content[:5] == b"%PDF-"
    assert client.get("/api/export/BAD!!.pdf").status_code == 400


def test_companyfacts_keeps_only_needed_rows_and_reads_once(monkeypatch):
    import json
    calls = []
    big = {**GAAP, "AccountsPayableCurrent": {"units": {"USD": [fact(1e9, "2024-01-01", "2024-12-31", "x", "2025-02-01")]}}}
    monkeypatch.setattr(sec_core, "sec_fetch", lambda url: calls.append(url) or json.dumps({"facts": {"us-gaap": big}}).encode())
    sec_core._facts.clear()
    first = sec_core.companyfacts("0000018230")
    assert sec_core.companyfacts("0000018230") is first and len(calls) == 1   # trends reuse the Pay Efficiency read
    assert "AccountsPayableCurrent" not in first                             # tags the app never reads are dropped
    rows = first["OperatingIncomeLoss"]["units"]["USD"]
    assert all(r["form"] == "10-K" and sec_core.days_apart(r["start"], r["end"]) >= 340 for r in rows)  # no quarters or 10-Qs
    # Same answers from the slim copy as from the full file
    assert sec_core.tenk_value(first, ["OperatingIncomeLoss"], "2024-12-31") == sec_core.tenk_value(GAAP, ["OperatingIncomeLoss"], "2024-12-31")
    for i in range(sec_core.FACTS_CACHE_SIZE + 1):
        sec_core.companyfacts(str(i))
    assert "0000018230" not in sec_core._facts and len(sec_core._facts) == sec_core.FACTS_CACHE_SIZE



def test_proxy_text_reads_primary_document_only(monkeypatch):
    import sec_core
    from types import SimpleNamespace
    body = "<html><body><p>" + "Our lead independent director chairs executive sessions. " * 1200 + "</p></body></html>"
    fetched = []
    monkeypatch.setattr(sec_core, "sec_fetch", lambda url: fetched.append(url) or body.encode())
    filing = SimpleNamespace(cik=18230, accession_no="0000018230-26-000010", primary_document="cat-def14a.htm", attachments=[])
    text, source = sec_core.get_proxy_text(filing)
    assert source == "primary document" and "lead independent director" in text
    assert fetched == ["https://www.sec.gov/Archives/edgar/data/18230/000001823026000010/cat-def14a.htm"]


def test_latest_filings_reads_full_history_only_when_recent_page_is_empty():
    import sec_core
    calls = []
    def get_filings(form, trigger_full_load=True):
        calls.append(trigger_full_load)
        return [] if not trigger_full_load else ["old proxy"]
    from types import SimpleNamespace
    assert sec_core.latest_filings(SimpleNamespace(get_filings=get_filings), "DEF 14A") == ["old proxy"]
    assert calls == [False, True]


def _slow(ticker):
    import time
    time.sleep(0.3)
    return ticker


def test_lookups_run_one_at_a_time():
    import threading, time
    import main
    started = time.time()
    threads = [threading.Thread(target=main.isolated, args=(_slow, t)) for t in ["CAT", "JPM", "GS"]]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert time.time() - started >= 0.9

def _killed(ticker):
    import os, signal
    os.kill(os.getpid(), signal.SIGKILL)


def _fails(ticker):
    raise ValueError(f"no filing for {ticker}")


def test_isolated_lookup_returns_result_and_reports_failures():
    import pytest
    import main
    assert main.isolated(lambda t: {"ticker": t, "rows": list(range(3))}, "CAT") == {"ticker": "CAT", "rows": [0, 1, 2]}
    with pytest.raises(RuntimeError, match="no filing for JPM"):
        main.isolated(_fails, "JPM")
    # A child stopped by the out-of-memory killer gets a plain explanation, and the server process carries on
    with pytest.raises(RuntimeError, match="too large for the free server's memory"):
        main.isolated(_killed, "C")


def _bundle(tmp_path, files):
    import io, json, tarfile
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as tar:
        for name, body in files.items():
            data = json.dumps(body).encode()
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tar.addfile(info, io.BytesIO(data))
    return buf.getvalue()


def fake_stream(blob):
    """Stands in for httpx.stream: a context manager whose response yields the blob."""
    from contextlib import contextmanager
    @contextmanager
    def stream(*a, **k):
        yield SimpleNamespace(raise_for_status=lambda: None, iter_bytes=lambda n: iter([blob[:10], blob[10:]]))
    return stream


def test_prebuilt_results_are_served_without_reading_sec(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    import main, prebuilt
    ready = {"ticker": "AAPL", "name": "Apple Inc.", "ratios": {"Gross Margin": "46.9%"}}
    blob = _bundle(tmp_path, {"manifest.json": {"built_at": "2026-10-20T06:00:00Z", "tickers": ["AAPL"]},
                              "AAPL/company.json": ready})
    monkeypatch.setattr(prebuilt, "DIR", str(tmp_path / "prebuilt"))
    monkeypatch.setattr(prebuilt.httpx, "stream", fake_stream(blob))
    cleared = []
    monkeypatch.setattr(prebuilt, "_on_refresh", [lambda: cleared.append(True)])
    assert prebuilt.refresh() and cleared
    assert prebuilt.status()["companies"] == 1

    def no_sec(t):
        raise AssertionError("read SEC for a prebuilt company")
    monkeypatch.setattr(main.sec_core, "get_financial_report", no_sec)
    main.cached_financials.cache_clear()
    body = TestClient(main.app).get("/api/company/aapl").json()
    assert body.pop("benchmark") is None and body == ready     # no benchmarks.json in this bundle
    # Anything not in the file is still read live
    assert prebuilt.get("governance", "AAPL") is None and prebuilt.get("company", "ZZZZ") is None
    main.cached_financials.cache_clear()


def test_new_build_is_picked_up_within_the_hour(monkeypatch):
    import prebuilt
    downloads, stamp = [], ["2026-10-20T22:10:00Z"]
    monkeypatch.setattr(prebuilt, "published_at", lambda: stamp[0])
    monkeypatch.setattr(prebuilt, "refresh", lambda: downloads.append(1) or prebuilt._status.update(loaded_at="x") or True)
    monkeypatch.setattr(prebuilt, "_status", {"loaded_at": None})
    monkeypatch.setattr(prebuilt, "_state", {"published": None, "refreshed": 0.0})
    prebuilt.check(now=100_000)                       # startup: downloads
    prebuilt.check(now=100_000 + 3600)                # same build an hour later: nothing to do
    assert len(downloads) == 1
    stamp[0] = "2026-10-20T22:35:00Z"                 # the weekly job published a new file
    prebuilt.check(now=100_000 + 7200)
    assert len(downloads) == 2
    prebuilt.check(now=100_000 + 7200 + 12 * 3600)    # 12 hours on, downloads anyway
    assert len(downloads) == 3


def test_prebuilt_download_failure_keeps_current_files(tmp_path, monkeypatch):
    import prebuilt
    monkeypatch.setattr(prebuilt, "DIR", str(tmp_path / "prebuilt"))
    def boom(*a, **k):
        raise OSError("network down")
    monkeypatch.setattr(prebuilt.httpx, "stream", boom)
    assert prebuilt.refresh() is False
    assert "network down" in prebuilt.status()["error"]


def test_prebuild_saves_each_lookup_and_skips_failures(tmp_path, monkeypatch):
    import json, os
    import prebuild
    monkeypatch.chdir(tmp_path)
    def fails(t):
        raise ValueError("no proxy")
    monkeypatch.setattr(prebuild, "LOOKUPS", {"company": lambda t: {"ticker": t}, "governance": fails})
    ticker, saved, failed = prebuild.build_one("MSFT")
    assert (ticker, saved) == ("MSFT", ["company"]) and "no proxy" in failed["governance"]
    assert json.load(open(os.path.join("prebuilt", "MSFT", "company.json"))) == {"ticker": "MSFT"}
    assert not os.path.exists(os.path.join("prebuilt", "MSFT", "governance.json"))
