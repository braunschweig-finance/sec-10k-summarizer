import io
import re
import time
import warnings
import pandas as pd
import streamlit as st
from docx import Document
from edgar import Company, set_identity
from google import genai

# Suppress internal warnings
warnings.filterwarnings("ignore")

# 1. SEC Compliance Identity
set_identity("Jacob Braunschweig jacob.braunschweig@gmail.com")

# 2. Capital IQ / NetAdvantage Style Layout
st.set_page_config(page_title="SEC Financial & Corporate Governance Terminal", layout="wide")
st.title("🏛️ SEC Financial & Corporate Governance Terminal")
st.caption("Direct SEC EDGAR Statement Extraction (XBRL), Deterministic Ratios & DEF 14A Governance Analysis")

# Sidebar Controls
st.sidebar.header("Terminal Navigation")
ticker = st.sidebar.text_input("Enter Ticker Symbol:", "CAT").upper().strip()
run_analysis = st.sidebar.button("Fetch & Analyze SEC Data", type="primary")

st.sidebar.markdown("---")
st.sidebar.markdown("### System Architecture (Prof. Guidelines)")
st.sidebar.markdown("""
- **Data Ingestion:** SEC EDGAR REST / XBRL
- **Financial Statements:** Form 10-K (in $ Millions)
- **Quantitative Engine:** Deterministic Python (0 LLM Tokens)
- **Governance Mining:** Form DEF 14A Proxy Statements
- **AI Synthesis Engine:** Google Gemini (Executive Memo)
- **Deliverables Export:** Formatted Word (.docx) & Plaintext (.txt)
""")

def is_date_or_period_column(col_name):
    """Detects whether a column represents an audited fiscal reporting period."""
    s = str(col_name).strip()
    if re.search(r"\b20[2-3][0-9]", s):
        return True
    if any(tag in s.upper() for tag in ["(FY)", "(CY)", "Q1", "Q2", "Q3", "Q4"]):
        return True
    return False

def latest_period_column(df):
    """Returns the most recent fiscal-year column (by end date), not just the first column."""
    period_cols = [c for c in df.columns if is_date_or_period_column(c)]
    if not period_cols:
        return None
    fy_cols = [c for c in period_cols if "(FY)" in str(c).upper()] or period_cols
    def end_date(c):
        m = re.search(r"\d{4}-\d{2}-\d{2}", str(c))
        return m.group(0) if m else ""
    return max(fy_cols, key=end_date)

def parse_num(val):
    """Converts an XBRL cell to float. Keeps zero, negatives and small values; returns None for blanks."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    if s.lower() in ["none", "nan", ""]:
        return None
    try:
        return float(s.replace(",", "").replace("$", "").replace("(", "-").replace(")", "").strip())
    except Exception:
        return None

def unit_kind(unit):
    """Classifies an XBRL unit name ("usd", "u_usd", "usdPerShare", "perShare", "shares", ...)."""
    u = str(unit).lower() if unit is not None and not pd.isna(unit) else "usd"
    if "share" in u:
        return "per_share" if "per" in u else "shares"
    if "usd" in u:
        return "usd"
    return "other"

def format_cell(val, unit):
    """Formats one statement value by its XBRL unit: USD in $ millions, per-share in $, shares in millions."""
    num = parse_num(val)
    if num is None:
        return ""
    kind = unit_kind(unit)
    if kind == "usd":
        return f"${num / 1_000_000:,.0f}"
    if kind == "per_share":
        return f"${num:,.2f}"
    if kind == "shares":
        return f"{num / 1_000_000:,.1f}M sh"
    return f"{num:,.2f}"

def clean_statement_df(df):
    """
    Cleans XBRL statements to match Capital IQ:
    - Retains row label and numeric fiscal date periods
    - Removes abstract category header rows that have no numeric values
    - Scales dollar values to millions ($M); per-share and share counts keep their own units
    """
    if df is None or df.empty:
        return None

    label_col = None
    for cand in ["label", "standard_concept", "concept"]:
        if cand in df.columns:
            label_col = cand
            break

    period_cols = [c for c in df.columns if is_date_or_period_column(c)]
    if not period_cols:
        return None

    work = df.copy()
    if "abstract" in work.columns:
        work = work[~work["abstract"].fillna(False).astype(bool)]

    # Drop header rows where every period value is blank or zero
    has_number = work[period_cols].apply(
        lambda row: any((parse_num(v) or 0) != 0 for v in row), axis=1
    )
    work = work[has_number]

    units = work["unit"] if "unit" in work.columns else pd.Series("usd", index=work.index)
    cleaned = pd.DataFrame(index=work.index)
    cleaned["Line Item"] = work[label_col] if label_col else work.index.astype(str)
    for col in period_cols:
        cleaned[col] = [format_cell(v, u) for v, u in zip(work[col], units)]
    return cleaned.set_index("Line Item")

def _norm_label(s):
    return str(s).lower().replace("’", "'").strip().rstrip(":").strip()

def extract_metric(df, concepts, labels):
    """
    Finds one consolidated line item in the latest fiscal year.
    Order: exact XBRL concepts in priority order, then exact row labels on company-specific
    (non us-gaap) concepts only, since labels on standard concepts can be misleading.
    Segment (dimensional), abstract and non-USD rows are skipped. Returns {"value", "concept", "label", "period"} or None.
    """
    if df is None or df.empty:
        return None
    period = latest_period_column(df)
    if period is None:
        return None

    rows = df
    if "dimension" in rows.columns:
        rows = rows[~rows["dimension"].fillna(False).astype(bool)]
    if "abstract" in rows.columns:
        rows = rows[~rows["abstract"].fillna(False).astype(bool)]
    if "unit" in rows.columns:
        rows = rows[rows["unit"].map(unit_kind) == "usd"]
    if rows.empty:
        return None

    # "us-gaap_Revenues" -> "Revenues"
    concept_names = rows["concept"].astype(str).str.split(r"[_:]", n=1).str[-1] if "concept" in rows.columns else None

    def first_hit(mask):
        for _, r in rows[mask].iterrows():
            n = parse_num(r[period])
            if n is not None:
                return {"value": n, "concept": str(r.get("concept", "")), "label": str(r.get("label", "")), "period": str(period)}
        return None

    if concept_names is not None:
        for c in concepts:
            hit = first_hit(concept_names == c)
            if hit:
                return hit
    if "label" in rows.columns:
        row_labels = rows["label"].map(_norm_label)
        custom = ~rows["concept"].astype(str).str.startswith("us-gaap") if "concept" in rows.columns else True
        for lbl in labels:
            hit = first_hit((row_labels == _norm_label(lbl)) & custom)
            if hit:
                return hit
    return None

def compute_ratios(raw_inc, raw_bal):
    """Pulls the ratio inputs from the latest 10-K and computes the four ratios. Returns (ratios, inputs)."""
    revenue = extract_metric(raw_inc, ["Revenues", "RevenuesNetOfInterestExpense"], [])
    if revenue is None:
        # Banks without a total revenue line: net interest income + noninterest income
        nii = extract_metric(raw_inc, ["InterestIncomeExpenseNet"], [])
        nonint = extract_metric(raw_inc, ["NoninterestIncome"], [])
        if nii and nonint:
            revenue = {"value": nii["value"] + nonint["value"], "concept": "Derived: Net interest income + Noninterest income",
                       "label": "Net interest income + Total noninterest income", "period": nii["period"]}
    if revenue is None:
        revenue = extract_metric(
            raw_inc,
            ["RevenueFromContractWithCustomerExcludingAssessedTax", "RevenueFromContractWithCustomerIncludingAssessedTax", "SalesRevenueNet"],
            ["Total sales and revenues", "Total net sales", "Total revenues", "Total net revenue", "Net sales", "Revenues", "Revenue"]
        )

    inputs = {
        "Revenue": revenue,
        "Cost of Revenue": extract_metric(
            raw_inc,
            ["CostOfGoodsAndServicesSold", "CostOfRevenue", "CostOfGoodsSold", "CostOfGoodsAndServiceExcludingDepreciationDepletionAndAmortization"],
            ["Cost of goods sold", "Cost of sales", "Cost of products sold", "Cost of revenue", "Total cost of revenue", "Total cost of sales"]
        ),
        "Gross Profit": extract_metric(raw_inc, ["GrossProfit"], ["Gross profit", "Gross margin"]),
        "Operating Income": extract_metric(
            raw_inc,
            ["OperatingIncomeLoss"],
            ["Operating profit", "Operating income", "Operating income (loss)", "Income from operations"]
        ),
        "Current Assets": extract_metric(raw_bal, ["AssetsCurrent"], ["Total current assets"]),
        "Current Liabilities": extract_metric(raw_bal, ["LiabilitiesCurrent"], ["Total current liabilities"]),
        "Total Liabilities": extract_metric(raw_bal, ["Liabilities"], ["Total liabilities"]),
        "Shareholders' Equity": extract_metric(
            raw_bal,
            ["StockholdersEquity", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"],
            ["Total shareholders' equity", "Total stockholders' equity", "Total equity"]
        ),
    }

    # Derived inputs are labeled as such so every number still traces to filed line items
    rev, cogs = inputs["Revenue"], inputs["Cost of Revenue"]
    if inputs["Gross Profit"] is None and rev and cogs:
        inputs["Gross Profit"] = {"value": rev["value"] - cogs["value"], "concept": "Derived: Revenue - Cost of Revenue",
                                  "label": "Revenue - Cost of Revenue", "period": rev["period"]}
    if inputs["Total Liabilities"] is None:
        liab_and_eq = extract_metric(raw_bal, ["LiabilitiesAndStockholdersEquity"], ["Total liabilities and shareholders' equity", "Total liabilities and stockholders' equity", "Total liabilities and equity"])
        total_eq = extract_metric(
            raw_bal,
            ["StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest", "StockholdersEquity"],
            ["Total equity", "Total shareholders' equity", "Total stockholders' equity"]
        )
        if liab_and_eq and total_eq:
            inputs["Total Liabilities"] = {"value": liab_and_eq["value"] - total_eq["value"],
                                           "concept": "Derived: Total liabilities & equity - Total equity",
                                           "label": "Total liabilities & equity - Total equity", "period": liab_and_eq["period"]}

    def v(name):
        return inputs[name]["value"] if inputs[name] else None

    def pct(num, den):
        return f"{num / den * 100:.1f}%" if num is not None and den is not None and den > 0 else "N/A"

    def times(num, den):
        return f"{num / den:.2f}x" if num is not None and den is not None and den > 0 else "N/A"

    equity = v("Shareholders' Equity")
    ratios = {
        "Gross Margin": pct(v("Gross Profit"), v("Revenue")),
        "Operating Margin": pct(v("Operating Income"), v("Revenue")),
        "Current Ratio": times(v("Current Assets"), v("Current Liabilities")),
        # Negative equity makes the ratio meaningless, so say so instead of showing a negative multiple
        "Total Liabilities-to-Equity": "N/M (negative equity)" if equity is not None and equity < 0
                                       else times(v("Total Liabilities"), equity),
    }
    return ratios, inputs

def clean_governance_text(raw_text):
    """
    Cleans raw SEC proxy text (works for any filer):
    - Removes underscores, horizontal lines, dashes, and boilerplate rules
    - Removes running "Proxy Statement" page headers and mailing addresses
    - Collapses whitespace
    """
    if not raw_text:
        return ""
    text = raw_text
    # Drop invisible characters and unify dashes/apostrophes so keyword searches match
    text = re.sub(r"[\u00ad\u200b\u200c\u200d\ufeff]", "", text)
    text = re.sub(r"[\u2010\u2011\u2012\u2013]", "-", text)
    text = text.replace("\u2019", "'")
    # Remove lines of underscores, dashes, equals signs
    text = re.sub(r"_{2,}", " ", text)
    text = re.sub(r"-{3,}", " ", text)
    text = re.sub(r"={3,}", " ", text)
    # Remove running SEC headers/footers and addresses
    text = re.sub(r"\b\d{4}\s+Proxy Statement\s+\d{1,3}\b", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"Proxy Statement\s*\d{4}\s*\d{1,3}", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"c/o Corporate Secretary.{0,200}?\d{5}", " ", text, flags=re.IGNORECASE)
    # Collapse whitespace
    text = re.sub(r"\s+", " ", text).strip()
    return text

# Keyword checks for the governance tiles. A hit means the proxy discusses the topic; the matching
# sentence is shown as evidence so the reader can confirm what the filing actually says.
GOVERNANCE_CHECKS = {
    "Lead Independent Director": [r"lead independent director", r"\blead director\b", r"presiding (?:independent )?director"],
    "Clawback Policy": [r"clawback", r"(?:incentive )?compensation recovery policy"],
    "Anti-Hedging Policy": [r"anti-hedging", r"hedging (?:and|or) pledging", r"prohibit\w*[^.]{0,200}hedg"],
    "Say-on-Pay Vote": [r"advisory vote to approve (?:the )?(?:compensation of (?:our )?named executive officers|(?:named )?executive (?:officer )?compensation)",
                        r"say[- ]on[- ]pay"],
}

def html_to_plain_text(html):
    """Strips tags from an HTML document, keeping block breaks as newlines."""
    from bs4 import BeautifulSoup
    if isinstance(html, bytes):
        html = html.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    # Inline XBRL proxies hide header facts in ix:header; drop them so they don't pollute the text
    for tag in soup.find_all(["ix:header"]):
        tag.decompose()
    return soup.get_text("\n")

def get_proxy_text(filing):
    """
    Returns (text, source) for a DEF 14A. Tries edgartools' text() first; if that fails or comes back
    too short to be a full proxy, reads every HTML document in the filing directly, since some
    proxies are split across several files or don't parse with the default extractor.
    """
    text, source = "", "none"
    try:
        text = filing.text() or ""
        source = "primary document"
    except Exception:
        text = ""
    if len(text) >= 50000:
        return text, source

    parts = []
    try:
        for att in filing.attachments:
            if not att.is_html() or att.is_report() or att.empty:
                continue
            content = att.download()
            if content:
                parts.append(html_to_plain_text(content))
    except Exception:
        pass
    combined = "\n\n".join(parts)
    if len(combined) > len(text):
        return combined, f"{len(parts)} HTML document(s) in the filing"
    return text, source

def find_evidence(text, patterns):
    """Returns the sentence around the first pattern match in the proxy, or None."""
    for p in patterns:
        m = re.search(p, text, flags=re.IGNORECASE)
        if m:
            start = max(text.rfind(". ", 0, m.start()) + 2, m.start() - 300, 0)
            end = text.find(". ", m.end())
            end = min(end + 1 if end != -1 else len(text), m.end() + 300)
            return text[start:end].strip()
    return None

def build_word_memo(ticker_sym, company_nm, ratios_dict, gov_dict, memo_body):
    """Generates a structured Word (.docx) Institutional Research Memo."""
    doc = Document()
    doc.add_heading(f"Institutional Research & Corporate Governance Brief", level=0)
    doc.add_paragraph(f"Company: {company_nm} ({ticker_sym}) | Source: Audited SEC EDGAR 10-K & DEF 14A")

    doc.add_heading("1. Executive Summary & Audited Ratios", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Financial Ratio"
    hdr_cells[1].text = "Audited Metric Value"
    for metric_name, val in ratios_dict.items():
        row_cells = table.add_row().cells
        row_cells[0].text = metric_name
        row_cells[1].text = str(val)

    doc.add_heading("2. Qualitative Corporate Governance Disclosures (DEF 14A)", level=1)
    for section_title, section_content in gov_dict.items():
        doc.add_heading(section_title, level=2)
        doc.add_paragraph(section_content[:1500] + ("..." if len(section_content) > 1500 else ""))

    doc.add_heading("3. Institutional Equity Research Synthesis", level=1)
    doc.add_paragraph(memo_body)

    file_buffer = io.BytesIO()
    doc.save(file_buffer)
    return file_buffer.getvalue()

def statement_frames(statement):
    """Returns (display_df, raw_df) for one statement, or (None, None) if the filing lacks it."""
    if statement is None:
        return None, None
    # standard=False keeps the filer's own line-item labels instead of edgartools' standardized ones
    display_df = statement.to_dataframe(view="standard", include_unit=True, standard=False)
    raw_df = statement.to_dataframe(presentation=False, include_unit=True, standard=False)
    return display_df, raw_df

@st.cache_data(show_spinner=False)
def get_sec_data(ticker_symbol):
    # Errors are raised, not returned, so a temporary SEC outage is not cached for the session
    company = Company(ticker_symbol)

    # 1. 10-K Audited Financial Statements
    financials = company.get_financials()
    disp_income, raw_income = statement_frames(financials.income_statement() if financials else None)
    disp_balance, raw_balance = statement_frames(financials.balance_sheet() if financials else None)
    disp_cashflow, _ = statement_frames(financials.cash_flow_statement() if financials else None)

    # 2. Form DEF 14A (Proxy Statement)
    proxy_filings = company.get_filings(form="DEF 14A")
    tenk_filings = company.get_filings(form="10-K")

    governance_sections = {}
    governance_checks = {}
    proxy_chars = 0
    proxy_source = ""
    proxy_url = ""
    proxy_date = ""
    tenk_url = tenk_filings[0].url if tenk_filings else ""

    if proxy_filings:
        latest_proxy = proxy_filings[0]
        proxy_url = latest_proxy.url
        proxy_date = str(latest_proxy.filing_date)
        proxy_raw, proxy_source = get_proxy_text(latest_proxy)
        proxy_chars = len(proxy_raw)

        proxy_flat = clean_governance_text(proxy_raw)
        for check, patterns in GOVERNANCE_CHECKS.items():
            governance_checks[check] = find_evidence(proxy_flat, patterns)

        proxy_lower = proxy_flat.lower()

        # Prefer matches after the first 15,000 characters to bypass the table of contents and cover page
        body_start = 15000 if len(proxy_flat) > 25000 else 0

        section_targets = {
            "Board Leadership & Lead Independent Director Structure": [
                "board leadership structure", "lead independent director", "director independence"
            ],
            "Executive Compensation & Pay-for-Performance (CD&A)": [
                "compensation discussion and analysis", "executive compensation program",
                "compensation philosophy", "summary compensation table"
            ],
            "Annual Meeting Agenda & Shareholder Voting Items": [
                "matters to be voted on", "proposals to be voted on", "election of directors",
                "shareholder proposals", "stockholder proposals"
            ],
            "Clawback Policies, Anti-Hedging & Risk Oversight": [
                "clawback policy", "compensation recovery policy",
                "anti-hedging", "hedging and pledging", "risk oversight"
            ]
        }

        for section_title, keywords in section_targets.items():
            found_pos = -1
            for kw in keywords:
                pos = proxy_lower.find(kw, body_start)
                if pos == -1:
                    pos = proxy_lower.find(kw)
                if pos != -1:
                    found_pos = pos
                    break
            if found_pos != -1:
                governance_sections[section_title] = proxy_flat[found_pos : found_pos + 3000]
            else:
                governance_sections[section_title] = "No matching section heading was found in the proxy text."

    return {
        "name": company.name,
        "income": clean_statement_df(disp_income),
        "balance": clean_statement_df(disp_balance),
        "cashflow": clean_statement_df(disp_cashflow),
        "raw_income": raw_income,
        "raw_balance": raw_balance,
        "gov_sections": governance_sections,
        "gov_checks": governance_checks,
        "proxy_url": proxy_url,
        "proxy_date": proxy_date,
        "proxy_chars": proxy_chars,
        "proxy_source": proxy_source,
        "tenk_url": tenk_url
    }

if ticker:
    with st.spinner(f"Extracting SEC 10-K statements and DEF 14A proxy data for {ticker}..."):
        try:
            data, err = get_sec_data(ticker), None
        except Exception as e:
            data, err = None, str(e)

    if err:
        st.error(f"Error fetching data: {err}")
    elif data:
        st.subheader(f"{data['name']} ({ticker})")

        # Capital IQ style filing access header
        col_meta1, col_meta2 = st.columns(2)
        if data["tenk_url"]:
            col_meta1.markdown(f"📄 [SEC EDGAR Form 10-K Annual Filing]({data['tenk_url']})")
        if data["proxy_url"]:
            col_meta2.markdown(f"🗳️ [SEC EDGAR Form DEF 14A Proxy Statement]({data['proxy_url']})")

        tab_stmt, tab_ratios, tab_gov, tab_memo = st.tabs([
            "📋 Financial Statements (10-K)",
            "📈 Financial Ratios",
            "🗳️ Corporate Governance (DEF 14A)",
            "🤖 Institutional Research Memo"
        ])

        # TAB 1: FINANCIAL STATEMENTS (10-K)
        with tab_stmt:
            st.markdown("### Audited Financial Statements (Direct SEC XBRL)")
            st.caption("Dollar figures in **\\$ Millions of USD**; per-share amounts in \\$; share counts in millions. Zero LLM tokens consumed.")

            stmt_view = st.radio(
                "Select Statement to Inspect:",
                ["Income Statement", "Balance Sheet", "Cash Flow"],
                horizontal=True
            )

            selected_df = None
            if stmt_view == "Income Statement":
                selected_df = data["income"]
            elif stmt_view == "Balance Sheet":
                selected_df = data["balance"]
            elif stmt_view == "Cash Flow":
                selected_df = data["cashflow"]

            if selected_df is not None and not selected_df.empty:
                st.dataframe(selected_df, width="stretch")
            else:
                st.info("Direct XBRL table not available for this statement.")

        # TAB 2: FINANCIAL RATIOS (Deterministic Python)
        with tab_ratios:
            st.markdown("### Deterministic Ratio Analysis")
            st.caption("Computed via Python from audited consolidated line items in the latest fiscal year (zero AI token consumption).")

            ratios_summary, ratio_inputs = compute_ratios(data.get("raw_income"), data.get("raw_balance"))
            gm = ratios_summary["Gross Margin"]
            om = ratios_summary["Operating Margin"]
            cr = ratios_summary["Current Ratio"]
            de = ratios_summary["Total Liabilities-to-Equity"]

            col1, col2, col3, col4 = st.columns(4)
            col1.metric("Gross Margin", gm)
            col2.metric("Operating Margin", om)
            col3.metric("Current Ratio", cr)
            col4.metric("Total Liabilities-to-Equity", de, help="Total liabilities / shareholders' equity. Includes all liabilities, not only debt.")

            st.caption("N/A means the filing does not report that line item on the face of its statements "
                       "(e.g., banks report no gross profit or current assets; some companies report no operating income).")

            st.markdown("#### Source Line Items")
            source_rows = []
            for name, hit in ratio_inputs.items():
                source_rows.append({
                    "Input": name,
                    "Value ($M)": f"{hit['value'] / 1_000_000:,.0f}" if hit else "Not reported",
                    "Filed Line Item": hit["label"] if hit else "",
                    "XBRL Concept": hit["concept"] if hit else "",
                    "Period": hit["period"] if hit else "",
                })
            st.dataframe(pd.DataFrame(source_rows).set_index("Input"), width="stretch")

        # TAB 3: CORPORATE GOVERNANCE (DEF 14A)
        with tab_gov:
            st.markdown("### Institutional Corporate Governance Profile")
            st.caption("Extracted directly from Form DEF 14A Proxy Statements. Structured per S&P Capital IQ standards.")

            if data["gov_sections"]:
                # Key Governance Indicators: keyword evidence from the proxy, not assumed values
                st.caption(f"Latest DEF 14A filed {data['proxy_date']} ({data['proxy_chars']:,} characters of text read from "
                           f"{data['proxy_source']}). Each tile shows whether the proxy discusses the topic; "
                           "read the quoted evidence below to confirm what it says.")
                if data["proxy_chars"] < 20000:
                    st.warning("Very little text could be read from this proxy statement, so the results below may be incomplete. "
                               "Open the DEF 14A link above to check the filing directly.")
                gov_cols = st.columns(len(GOVERNANCE_CHECKS))
                for col, (check, evidence) in zip(gov_cols, data["gov_checks"].items()):
                    col.metric(check, "Discussed" if evidence else "Not found")
                with st.expander("Evidence from the proxy statement"):
                    for check, evidence in data["gov_checks"].items():
                        st.markdown(f"**{check}:** " + (f"\"{evidence}\"" if evidence else "No matching text found."))

                st.markdown("---")

                # Clean Governance Modules
                for heading, text_body in data["gov_sections"].items():
                    with st.container(border=True):
                        st.subheader(f"📑 {heading}")
                        # Filter out stray single characters or short noise lines
                        clean_paragraphs = [p for p in text_body.split(". ") if len(p.strip()) > 20]
                        formatted_text = ". ".join(clean_paragraphs[:12])
                        st.write(formatted_text + ("." if formatted_text and not formatted_text.endswith(".") else ""))
            else:
                st.warning("No DEF 14A proxy filing located for this ticker.")

        # TAB 4: AI RESEARCH MEMO (Institutional Synthesis)
        with tab_memo:
            st.markdown("### Institutional Research Memo Synthesis")
            st.caption("Synthesizes quantitative 10-K statement metrics and qualitative DEF 14A proxy governance into an executive report.")

            # The memo is kept in session state so download clicks (which rerun the script) don't erase it
            # and don't spend another Gemini request.
            memo_key = f"memo_{ticker}"

            if st.button("Generate Institutional Research Memo", type="primary"):
                with st.spinner("Synthesizing 10-K & DEF 14A disclosures with Gemini..."):
                    models_to_try = ["gemini-3.8-flash", "gemini-3.6-flash"]
                    response = None
                    last_error = None

                    governance_text = "\n\n".join(f"{k}:\n{v}" for k, v in data["gov_sections"].items())
                    governance_flags = "\n".join(
                        f"- {k}: {'discussed' if v else 'not found'}" for k, v in data["gov_checks"].items()
                    )

                    memo_prompt = f"""
                    You are a senior equity research analyst preparing an institutional investment and corporate governance memo.
                    Company: {data['name']} ({ticker})

                    AUDITED 10-K QUANTITATIVE METRICS:
                    - Gross Margin: {gm}
                    - Operating Margin: {om}
                    - Current Ratio: {cr}
                    - Balance Sheet Leverage (Total Liabilities / Equity): {de}

                    DEF 14A PROXY CORPORATE GOVERNANCE DISCLOSURES:
                    {governance_flags}

                    {governance_text[:6000]}

                    DELIVERABLES:
                    Structure your memo in clean Markdown matching an institutional investment committee brief.
                    Use only the figures and disclosures above; if something is not provided, say it was not disclosed.
                    1. Executive Summary & Financial Quality (Evaluate core profitability, operating margin resilience, liquidity, and balance sheet leverage).
                    2. Corporate Governance Evaluation (Analyze executive compensation alignment, CEO-to-median employee pay ratio, board of directors committee independence, and lead independent director safeguards).
                    3. Annual Meeting & Shareholder Proposals (Review proposals submitted for shareholder vote, management recommendations, and shareholder rights).
                    4. Strategic Capital Allocation & Risk Verdict (Assess CapEx, share repurchases, dividend coverage, and governance risk score).
                    """

                    try:
                        api_key = st.secrets["GEMINI_API_KEY"]
                    except Exception:
                        api_key = None

                    if not api_key:
                        last_error = "GEMINI_API_KEY is not set in the app's secrets."
                    else:
                        client = genai.Client(api_key=api_key)
                        # At most 2 requests per model. A 429 (rate limit) moves straight to the next
                        # model instead of retrying, since retries count against the per-minute quota.
                        for model_name in models_to_try:
                            for attempt in range(2):
                                try:
                                    response = client.models.generate_content(
                                        model=model_name,
                                        contents=memo_prompt
                                    )
                                    break
                                except Exception as err:
                                    last_error = err
                                    if attempt == 0 and ("503" in str(err) or "UNAVAILABLE" in str(err)):
                                        time.sleep(3)
                                        continue
                                    break
                            if response:
                                break

                    if not response:
                        if "429" in str(last_error) or "RESOURCE_EXHAUSTED" in str(last_error):
                            st.error("Gemini rate limit reached (free-tier requests-per-minute cap). Wait a minute and try again.")
                        else:
                            st.error(f"Failed to generate memo: {last_error}")
                    else:
                        usage = None
                        if getattr(response, "usage_metadata", None):
                            usage = (response.usage_metadata.prompt_token_count, response.usage_metadata.candidates_token_count)
                        st.session_state[memo_key] = {"text": response.text, "usage": usage}

            memo = st.session_state.get(memo_key)
            if memo:
                st.markdown(memo["text"])
                if memo["usage"]:
                    st.info(f"Tokens Consumed — Input: {memo['usage'][0]} | Output: {memo['usage'][1]}")

                col_dl1, col_dl2 = st.columns(2)

                col_dl1.download_button(
                    label="📥 Download Research Memo (.txt)",
                    data=memo["text"],
                    file_name=f"{ticker}_Institutional_Memo.txt",
                    mime="text/plain"
                )

                docx_bytes = build_word_memo(
                    ticker_sym=ticker,
                    company_nm=data['name'],
                    ratios_dict=ratios_summary,
                    gov_dict=data['gov_sections'],
                    memo_body=memo["text"]
                )
                col_dl2.download_button(
                    label="📄 Download Institutional Memo (.docx)",
                    data=docx_bytes,
                    file_name=f"{ticker}_Institutional_Research_Brief.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
